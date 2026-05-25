"""
Vercel Python Serverless Function — 生成带 Word 批注的修改版 DOCX

POST /api/generate-docx
Body: { "docxBase64": "...", "optimizations": [...], "jobTitle": "..." }
Response: DOCX 二进制流 (application/vnd.openxmlformats-officedocument.wordprocessingml.document)
"""

import base64
import io
import json
import sys
import traceback
from datetime import datetime
from http.server import BaseHTTPRequestHandler

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn, nsmap
from lxml import etree


# ── Word 批注（Comment）支持 ──────────────────────────────────────────────

class CommentsPart:
    """管理 Word 文档的 comments.xml part，支持添加批注。"""

    def __init__(self, document):
        self.document = document
        self.comment_id = 0
        self._comments_element = None
        self._comments_part = None
        self._init_comments()

    def _init_comments(self):
        """获取或创建 comments part 和 comments 根元素。"""
        doc_part = self.document.part

        # 尝试获取已有的 comments part
        try:
            comments_rels = doc_part.rels._rels_by_rtype.get(
                'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments'
            )
            if comments_rels:
                rel = list(comments_rels.values())[0]
                self._comments_part = rel.target_part
                self._comments_element = etree.fromstring(self._comments_part.blob)
                # 找到最大 comment id
                existing_ids = self._comments_element.findall('.//' + qn('w:comment'))
                for c in existing_ids:
                    cid = int(c.get(qn('w:id'), 0))
                    if cid >= self.comment_id:
                        self.comment_id = cid + 1
                return
        except Exception:
            pass

        # 创建新的 comments.xml part
        comments_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:comments xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
            'xmlns:mo="http://schemas.microsoft.com/office/mac/office/2008/main" '
            'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:mv="urn:schemas-microsoft-com:mac:vml" '
            'xmlns:o="urn:schemas-microsoft-com:office:office" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:w10="urn:schemas-microsoft-com:office:word" '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:wne="http://schemas.microsoft.com/office/word/2006/wordml" '
            'xmlns:sl="http://schemas.openxmlformats.org/schemaLibrary/2006/main">'
            '</w:comments>'
        )
        self._comments_element = etree.fromstring(comments_xml.encode('utf-8'))

        content_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml'
        partname = '/word/comments.xml'

        # 使用 rels 添加 part
        doc_part.relate_to(
            partname,
            'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments',
            is_external=False
        )

        # 获取新创建的 part
        for rel in doc_part.rels.values():
            if rel.reltype == 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments':
                self._comments_part = rel.target_part
                break

        if self._comments_part is None:
            raise RuntimeError('Failed to create comments part')

        self._comments_part._blob = etree.tostring(
            self._comments_element, xml_declaration=True, encoding='UTF-8', standalone=True
        )
        self._comments_part.content_type = content_type

    def add_comment(self, text, author='Resume Optimizer'):
        """添加一条批注，返回批注 ID。"""
        cid = self.comment_id
        self.comment_id += 1

        comment_el = etree.SubElement(
            self._comments_element,
            qn('w:comment'),
            {
                qn('w:id'): str(cid),
                qn('w:author'): author,
                qn('w:date'): datetime.utcnow().isoformat() + 'Z',
            }
        )

        # 为每个段落添加 w:p
        for paragraph_text in text.split('\n'):
            p_el = etree.SubElement(comment_el, qn('w:p'))
            r_el = etree.SubElement(p_el, qn('w:r'))
            rpr_el = etree.SubElement(r_el, qn('w:rPr'))
            t_el = etree.SubElement(r_el, qn('w:t'))
            t_el.text = paragraph_text
            t_el.set(qn('xml:space'), 'preserve')

        self._save_comments()
        return cid

    def _save_comments(self):
        """将 comments XML 写回 part blob。"""
        if self._comments_part is None:
            raise RuntimeError('Comments part not initialized')

        self._comments_part._blob = etree.tostring(
            self._comments_element, xml_declaration=True, encoding='UTF-8', standalone=True
        )


def add_comment_to_paragraph(paragraph, comment_id):
    """在段落中插入 commentRangeStart、commentRangeEnd 和 commentReference。"""
    p_element = paragraph._element

    # commentRangeStart
    start = etree.Element(qn('w:commentRangeStart'), {qn('w:id'): str(comment_id)})
    # commentRangeEnd
    end = etree.Element(qn('w:commentRangeEnd'), {qn('w:id'): str(comment_id)})
    # commentReference
    ref = etree.Element(
        qn('w:r'),
        {
            qn('w:id'): str(comment_id),
        }
    )
    ref_r = etree.SubElement(ref, qn('w:rPr'))
    ref_comment = etree.SubElement(ref_r, qn('w:commentReference'), {qn('w:id'): str(comment_id)})

    # 在第一个 run 之前插入 commentRangeStart
    first_run = p_element.find(qn('w:r'))
    if first_run is not None:
        p_element.insert(list(p_element).index(first_run), start)
    else:
        p_element.append(start)

    # 在最后一个 run 之后插入 commentRangeEnd 和 commentReference
    p_element.append(end)
    p_element.append(ref)


def build_comment_text(opt):
    """根据 optimization 条目构建详细批注文本。"""
    parts = []
    mod_type = opt.get('modification_type', '未分类')
    old = opt.get('old_text', '')
    new = opt.get('new_text', '')
    comment = opt.get('comment', '')

    parts.append(f'【修改类型】{mod_type}')
    if old:
        old_brief = old[:100] + ('...' if len(old) > 100 else '')
        parts.append(f'【原文】{old_brief}')
    if new:
        new_brief = new[:100] + ('...' if len(new) > 100 else '')
        parts.append(f'【修改后】{new_brief}')
    elif new == '' and old:
        parts.append('【修改后】（已删除）')
    if comment:
        parts.append(f'【详细说明】{comment}')

    return '\n'.join(parts)


def apply_optimizations(doc, optimizations):
    """遍历 optimizations，在文档中匹配并替换文本，添加批注。

    Returns:
        dict: { "applied": int, "skipped": int, "total_comments": int }
    """
    comments = CommentsPart(doc)
    applied = 0
    skipped = 0
    total_comments = 0

    for opt in optimizations:
        old_text = (opt.get('old_text') or '').strip()
        new_text = opt.get('new_text', '')
        if not old_text:
            skipped += 1
            continue

        matched = False

        # 遍历所有段落进行匹配
        for para in doc.paragraphs:
            para_text = para.text.strip()
            if not para_text:
                continue

            # 精确匹配
            if old_text in para.text:
                matched = True
                _replace_in_paragraph(para, old_text, new_text, opt, comments)
                applied += 1
                total_comments += 1
                break

            # 去空格模糊匹配
            normalized_old = ''.join(old_text.split())
            normalized_para = ''.join(para.text.split())
            if normalized_old and normalized_old in normalized_para:
                matched = True
                _replace_in_paragraph(para, old_text, new_text, opt, comments)
                applied += 1
                total_comments += 1
                break

        if not matched:
            # 尝试在表格中匹配
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if old_text in para.text:
                                matched = True
                                _replace_in_paragraph(para, old_text, new_text, opt, comments)
                                applied += 1
                                total_comments += 1
                                break
                            normalized_old = ''.join(old_text.split())
                            normalized_para = ''.join(para.text.split())
                            if normalized_old and normalized_old in normalized_para:
                                matched = True
                                _replace_in_paragraph(para, old_text, new_text, opt, comments)
                                applied += 1
                                total_comments += 1
                                break
                        if matched:
                            break
                    if matched:
                        break
                if matched:
                    break

        if not matched:
            skipped += 1

    return {
        'applied': applied,
        'skipped': skipped,
        'total_comments': total_comments,
    }


def _replace_in_paragraph(para, old_text, new_text, opt, comments):
    """替换段落中的文本并添加批注。"""
    # 构建批注文本
    comment_text = build_comment_text(opt)

    # 如果整个段落就是 old_text，替换更彻底
    if para.text.strip() == old_text.strip():
        # 清空段落并写入新文本
        for run in para.runs:
            run.text = ''
        if para.runs:
            para.runs[0].text = new_text
        else:
            para.add_run(new_text)
    else:
        # 部分替换：在 runs 中查找并替换
        _replace_text_in_runs(para, old_text, new_text)

    # 添加 Word 批注
    if new_text:
        cid = comments.add_comment(comment_text, '简历优化助手')
    else:
        cid = comments.add_comment(
            comment_text + '\n【修改后】（已删除）', '简历优化助手'
        )
    add_comment_to_paragraph(para, cid)


def _replace_text_in_runs(para, old_text, new_text):
    """在段落的 runs 中查找并替换文本。"""
    remaining = old_text
    for run in para.runs:
        if not remaining:
            break
        if remaining in run.text:
            run.text = run.text.replace(remaining, new_text, 1)
            remaining = ''
        elif run.text and run.text in remaining:
            run.text = ''
            remaining = remaining.replace(run.text, '', 1)
        else:
            # 跨 run 的部分匹配
            overlap = _find_overlap(run.text, remaining)
            if overlap:
                run.text = run.text[:run.text.rfind(overlap)] + new_text
                remaining = ''


def _find_overlap(text_a, text_b):
    """找到 text_a 后缀与 text_b 前缀的最长重叠。"""
    max_len = min(len(text_a), len(text_b))
    for length in range(max_len, 0, -1):
        if text_a.endswith(text_b[:length]):
            return text_b[:length]
    return ''


# ── CORS 工具 ─────────────────────────────────────────────────────────────

def add_cors_headers(handler):
    handler.send_header('Access-Control-Allow-Origin', '*')
    handler.send_header('Access-Control-Allow-Methods', 'POST, OPTIONS')
    handler.send_header('Access-Control-Allow-Headers', 'Content-Type')


# ── HTTP Request Handler ───────────────────────────────────────────────────

class handler(BaseHTTPRequestHandler):
    """Vercel Python Serverless Function — 入口点"""

    def do_OPTIONS(self):
        self.send_response(204)
        add_cors_headers(self)
        self.end_headers()

    def do_POST(self):
        try:
            # 读取请求 body
            content_length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(content_length) if content_length > 0 else b''

            if not body:
                self._send_json_error(400, '请求 body 为空')
                return

            try:
                data = json.loads(body)
            except json.JSONDecodeError:
                self._send_json_error(400, 'JSON 解析失败，请确认请求格式正确')
                return

            # 校验必填字段
            docx_base64 = data.get('docxBase64', '')
            optimizations = data.get('optimizations', [])
            job_title = data.get('jobTitle', '未指定职位')

            if not docx_base64:
                self._send_json_error(400, '缺少 docxBase64 字段')
                return
            if not isinstance(optimizations, list) or len(optimizations) == 0:
                self._send_json_error(400, 'optimizations 必须是非空数组')
                return

            # 解码 base64 → DOCX
            try:
                docx_bytes = base64.b64decode(docx_base64)
            except Exception as e:
                self._send_json_error(400, f'base64 解码失败: {str(e)}')
                return

            # 用 python-docx 打开
            try:
                doc = Document(io.BytesIO(docx_bytes))
            except Exception as e:
                self._send_json_error(400, f'无法打开 DOCX 文件: {str(e)}')
                return

            # 应用优化
            stats = apply_optimizations(doc, optimizations)

            # 保存到内存
            output_buffer = io.BytesIO()
            doc.save(output_buffer)
            output_buffer.seek(0)
            output_bytes = output_buffer.getvalue()

            # 返回二进制 DOCX
            self.send_response(200)
            add_cors_headers(self)
            self.send_header('Content-Type',
                             'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            self.send_header('Content-Length', str(len(output_bytes)))
            self.send_header('X-Optimization-Stats',
                             f"applied={stats['applied']},skipped={stats['skipped']},comments={stats['total_comments']}")
            self.end_headers()
            self.wfile.write(output_bytes)

        except Exception as e:
            traceback.print_exc()
            self._send_json_error(500, f'服务器内部错误: {str(e)}')

    def _send_json_error(self, status_code, message):
        """发送 JSON 格式的错误响应。"""
        error_body = json.dumps({
            'error': 'DOCX Generation Error' if status_code >= 500 else 'Bad Request',
            'message': message,
            'errorMessage': message,
        }, ensure_ascii=False).encode('utf-8')

        try:
            self.send_response(status_code)
            add_cors_headers(self)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(error_body)))
            self.end_headers()
            self.wfile.write(error_body)
        except Exception:
            pass
