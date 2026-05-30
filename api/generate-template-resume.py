"""
Vercel Python Serverless Function — 模板化简历生成
POST /api/generate-template-resume
Body: { "templateBase64": "...", "resumeText": "...", "result": {...}, "jdText": "..." }
Response: { "docxBase64": "...", "message": "..." }
"""

import base64
import json
import re
from io import BytesIO
from http.server import BaseHTTPRequestHandler

from docx import Document
from docx.shared import RGBColor


def add_cors(handler):
    handler.send_header('Access-Control-Allow-Origin', '*')
    handler.send_header('Access-Control-Allow-Methods', 'POST, OPTIONS')
    handler.send_header('Access-Control-Allow-Headers', 'Content-Type')


class handler(BaseHTTPRequestHandler):

    def do_OPTIONS(self):
        self.send_response(204)
        add_cors(self)
        self.end_headers()

    def do_POST(self):
        try:
            content_length = int(self.headers.get('Content-Length', 0))
            body = self.rfile.read(content_length) if content_length > 0 else b''
            data = json.loads(body)

            template_b64 = data.get('templateBase64', '')
            resume_text = data.get('resumeText', '')
            result = data.get('result', {})
            jd_text = data.get('jdText', '')

            if not template_b64:
                self._json_error(400, '缺少模板文件')
                return

            # 解码模板
            template_bytes = base64.b64decode(template_b64)
            doc = Document(BytesIO(template_bytes))

            # 提取数据
            display = result.get('display', {})
            diagnostic = result.get('diagnosticReport', {})

            # 填充模板
            fill_resume(doc, resume_text, display, diagnostic, jd_text)

            # 保存到内存
            output = BytesIO()
            doc.save(output)
            output.seek(0)
            out_b64 = base64.b64encode(output.read()).decode()

            self.send_response(200)
            add_cors(self)
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(json.dumps({'docxBase64': out_b64, 'message': '生成成功'}).encode())

        except Exception as e:
            import traceback
            traceback.print_exc()
            self._json_error(500, f'生成失败: {str(e)}')

    def _json_error(self, status_code, message):
        self.send_response(status_code)
        add_cors(self)
        self.send_header('Content-Type', 'application/json')
        self.end_headers()
        self.wfile.write(json.dumps({'message': message}).encode())


# ========== 核心填充逻辑 ==========

def fill_resume(doc, resume_text, display, diagnostic, jd_text):
    paragraphs = doc.paragraphs
    resume_info = parse_resume_text(resume_text)
    experiences = display.get('experienceShowcase', [])

    # P0: 姓名
    name = resume_info.get('name', '')
    if name and len(paragraphs) > 0:
        replace_run_text(paragraphs[0], '姓名', name)

    # P1: 基本信息
    if len(paragraphs) > 1:
        p1 = paragraphs[1]
        replace_run_text(p1, 'xxxx', resume_info.get('political', ''))
        replace_run_text(p1, 'xxxxx', resume_info.get('phone', ''))
        replace_run_text(p1, 'xxxxxx', resume_info.get('email', ''))

    # P3: 教育信息
    edu = resume_info.get('education', {})
    if len(paragraphs) > 3:
        p3 = paragraphs[3]
        replace_run_text(p3, 'xxxxxxxxxxx大学', edu.get('school', ''))
        replace_run_text(p3, 'xxxxxxxx专业', edu.get('major', ''))
        if edu.get('start_year'):
            replace_nth_run_text(p3, 'xx', edu['start_year'], 1)
        if edu.get('end_year'):
            replace_nth_run_text(p3, 'xx', edu['end_year'], 3)
        if edu.get('gpa'):
            replace_nth_run_text(p3, 'xx', edu['gpa'], 4)
        if edu.get('rank'):
            replace_nth_run_text(p3, 'xx', edu['rank'], 5)

    # P4: 主修课程
    if len(paragraphs) > 4:
        courses = edu.get('courses', '') or '待补充'
        replace_run_text(paragraphs[4], 'xxxxxxxx', courses)

    # P5: 荣誉奖项
    if len(paragraphs) > 5:
        honors = resume_info.get('honors', '')
        if honors:
            replace_paragraph_runs(paragraphs[5], f'荣誉奖项：{honors}')

    # P6-P9: 实习经历
    intern_exp = find_experience(experiences, '实习')
    if intern_exp and len(paragraphs) > 9:
        if intern_exp.get('titleLine'):
            replace_paragraph_runs(paragraphs[7], intern_exp['titleLine'])
        bullets = extract_bullets(intern_exp.get('fullVersion', ''))
        if len(bullets) > 0:
            fill_bullet_paragraph(paragraphs[8], bullets[0])
        if len(bullets) > 1:
            fill_bullet_paragraph(paragraphs[9], bullets[1])

    # P10-P13: 项目经历
    proj_exp = find_experience(experiences, '项目')
    if proj_exp and len(paragraphs) > 13:
        if proj_exp.get('titleLine'):
            replace_paragraph_runs(paragraphs[11], proj_exp['titleLine'])
        bullets = extract_bullets(proj_exp.get('fullVersion', ''))
        if len(bullets) > 0:
            fill_bullet_paragraph(paragraphs[12], bullets[0])
        if len(bullets) > 1:
            fill_bullet_paragraph(paragraphs[13], bullets[1])

    # P14-P16: 校园经历
    campus_exp = find_experience(experiences, '校园')
    if campus_exp and len(paragraphs) > 16:
        if campus_exp.get('titleLine'):
            replace_paragraph_runs(paragraphs[15], campus_exp['titleLine'])
        bullets = extract_bullets(campus_exp.get('fullVersion', ''))
        if len(bullets) > 0:
            fill_bullet_paragraph(paragraphs[16], bullets[0])

    # P18-P22: 技能/优势
    skills_info = resume_info.get('skills', {})

    if len(paragraphs) > 18 and skills_info.get('certificates'):
        replace_run_text(paragraphs[18], 'xxxxx', skills_info['certificates'])

    if len(paragraphs) > 19 and skills_info.get('competitions'):
        replace_run_text(paragraphs[19], 'xxxxx', skills_info['competitions'])

    if len(paragraphs) > 20:
        tech = skills_info.get('technical', '') or '待补充'
        replace_run_text(paragraphs[20], 'xxxxxx。', tech)

    if len(paragraphs) > 21:
        self_eval = skills_info.get('selfEval', '') or '待补充'
        replace_run_text(paragraphs[21], 'xxxxxx。', self_eval)

    if len(paragraphs) > 22:
        hobbies = skills_info.get('hobbies', '') or '球类运动、写作'
        replace_run_text(paragraphs[22], 'xxxxxx。', hobbies)

    # 量化标红
    highlight_quantification(doc)

    # 末尾附加区
    append_ai_section(doc, display, diagnostic)


# ========== 工具函数 ==========

def replace_run_text(paragraph, old_text, new_text):
    for run in paragraph.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return True
    return False


def replace_nth_run_text(paragraph, old_text, new_text, n):
    count = 0
    for run in paragraph.runs:
        if old_text in run.text:
            count += 1
            if count == n:
                run.text = run.text.replace(old_text, new_text, 1)
                return True
    return False


def replace_paragraph_runs(paragraph, new_text):
    if not paragraph.runs:
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ''


def fill_bullet_paragraph(paragraph, bullet_text):
    if '：' in bullet_text:
        parts = bullet_text.split('：', 1)
        label = parts[0].strip('*').strip()
        content = parts[1].strip()
        if paragraph.runs:
            paragraph.runs[0].text = label
            paragraph.runs[0].bold = True
            if len(paragraph.runs) > 1:
                paragraph.runs[1].text = '：'
                paragraph.runs[1].bold = True
            if len(paragraph.runs) > 2:
                paragraph.runs[2].text = content
                paragraph.runs[2].bold = False
    else:
        if paragraph.runs:
            paragraph.runs[0].text = bullet_text


def parse_resume_text(text):
    info = {}
    lines = [l.strip() for l in text.strip().split('\n') if l.strip()]
    if not lines:
        return info

    info['name'] = re.sub(r'^[姓\s名：:]+', '', lines[0]).strip()

    for line in lines[:8]:
        phone_m = re.search(r'1[3-9]\d{9}', line)
        if phone_m:
            info['phone'] = phone_m.group()
        email_m = re.search(r'[\w.\-]+@[\w.\-]+\.\w+', line)
        if email_m:
            info['email'] = email_m.group()
        if '党员' in line or '团员' in line or '群众' in line:
            info['political'] = re.search(r'(?:党员|团员|群众)', line).group()

    edu = {}
    for line in lines:
        if any(kw in line for kw in ['大学', '学院', '学校', '本科', '硕士', '博士']):
            years = re.findall(r'20\d{2}', line)
            if len(years) >= 2:
                edu['start_year'] = years[0]
                edu['end_year'] = years[1]
            major_m = re.search(r'([一-鿿]{2,15})(?:专业|系)', line)
            if major_m:
                edu['major'] = major_m.group(1) + '专业'
            gpa_m = re.search(r'GPA[：:]\s*(\d+\.?\d*)', line, re.IGNORECASE)
            if gpa_m:
                edu['gpa'] = gpa_m.group(1)
            rank_m = re.search(r'(?:top|前)\s*(\d+)%', line, re.IGNORECASE)
            if rank_m:
                edu['rank'] = rank_m.group(1)
            school_m = re.search(r'([一-鿿]{2,15}(?:大学|学院|学校))', line)
            if school_m:
                edu['school'] = school_m.group(1)
            break
    info['education'] = edu

    for line in lines:
        if '主修课程' in line or '核心课程' in line:
            edu['courses'] = re.sub(r'^.*?[：:]\s*', '', line)
        if '荣誉' in line or '奖学金' in line or '奖项' in line:
            info['honors'] = re.sub(r'^.*?[：:]\s*', '', line)

    skills = {}
    for line in lines:
        if '证书' in line or 'CET' in line:
            skills['certificates'] = re.sub(r'^.*?[：:]\s*', '', line)
        if '比赛' in line or '竞赛' in line:
            skills['competitions'] = re.sub(r'^.*?[：:]\s*', '', line)
        if '专业技能' in line or '技术栈' in line:
            skills['technical'] = re.sub(r'^.*?[：:]\s*', '', line).rstrip('。.')
        if '自我评价' in line or '自我介绍' in line:
            skills['selfEval'] = re.sub(r'^.*?[：:]\s*', '', line).rstrip('。.')
        if '爱好' in line or '兴趣' in line:
            skills['hobbies'] = re.sub(r'^.*?[：:]\s*', '', line).rstrip('。.')
    info['skills'] = skills

    return info


def find_experience(experiences, keyword):
    for exp in experiences:
        if keyword in exp.get('name', ''):
            return exp
    return experiences[0] if experiences else None


def extract_bullets(full_version):
    if not full_version:
        return []
    lines = re.split(r'\n', full_version)
    return [re.sub(r'^[\-•\d.]+\s*', '', l).strip() for l in lines if len(l.strip()) > 5]


def highlight_quantification(doc):
    """成果型量化数据标红"""
    for para in doc.paragraphs:
        for run in para.runs:
            text = run.text or ''
            if not text:
                continue
            # 跳过标题行
            if run.bold and run.font.size and run.font.size > 140000:
                continue
            # 匹配成果型数字
            if re.search(r'(提升|降低|减少|增加|节省|缩短|优化|管理|服务|覆盖|处理|交付|支撑|从.{1,8}至)\s*(?:约?\s*)?\d+', text):
                run.font.color = RGBColor(0xFF, 0x00, 0x00)
            elif re.search(r'\d+%', text) and not re.search(r'(GPA|top|前)\s*\d', text):
                run.font.color = RGBColor(0xFF, 0x00, 0x00)


def append_ai_section(doc, display, diagnostic):
    """末尾附加区"""
    doc.add_paragraph('')

    # 灰色分隔标题
    p_sep = doc.add_paragraph('')
    run_sep = p_sep.add_run('【以下为 AI 推断补充，供参考，请自行核实后决定是否采用】')
    run_sep.font.color = RGBColor(0x80, 0x80, 0x80)

    missing = display.get('missingInfoSuggestions', [])
    for item in missing:
        p = doc.add_paragraph('')
        run = p.add_run(f'• {item}')
        run.font.color = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph('')

    # 匹配度评分
    p_title = doc.add_paragraph('')
    run_t = p_title.add_run('优化后简历打分 · 匹配度计算（100 分制）')
    run_t.bold = True

    if diagnostic:
        dims = diagnostic.get('dimensions', {})
        for key, label in [('jdMatch', 'JD 匹配度'), ('quantification', '量化成果'),
                           ('structure', '结构与逻辑'), ('language', '语言专业度'), ('ats', 'ATS 友好度')]:
            d = dims.get(key, {})
            p = doc.add_paragraph('')
            run = p.add_run(f'{label}：{d.get("score", 0)}/{d.get("maxScore", 0)}')
            run.bold = True
            if d.get('detail'):
                p.add_run(f'  {d["detail"]}')

        p_ov = doc.add_paragraph('')
        run_ov = p_ov.add_run(f'综合得分：{diagnostic.get("overallScore", 0)}/100')
        run_ov.bold = True

    doc.add_paragraph('')
    p_hl = doc.add_paragraph('')
    p_hl.add_run('优化亮点：').bold = True
    for h in ['🔑 关键词优化', '📊 量化成果', '🎯 技能匹配', '✨ 措辞优化', '📐 结构调整', '🤖 ATS 优化']:
        doc.add_paragraph(f'  {h}')
